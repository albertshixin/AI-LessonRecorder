# -*- coding: utf-8 -*-
"""逐字稿 Word 导出：python-docx 直接由时间轴事件生成图文文档"""
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.core.timeline import TimelineEvent, SlideEvent
from src.export.paragraph import iter_grouped
from src.utils.timeutil import fmt_ts


def _style_doc(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def export_transcript_docx(session_root: Path, course_name: str,
                           events: list[TimelineEvent], include_images: bool = True) -> Path:
    """导出逐字稿 Word 文档，返回文件路径"""
    doc = Document()
    _style_doc(doc)

    title = doc.add_heading(f"课程逐字稿：{course_name}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sorted_events = sorted(events, key=lambda e: e.ts)
    has_slide = any(isinstance(e, SlideEvent) for e in sorted_events)
    opened = False  # 是否已输出第一个小节标题

    for kind, ev in iter_grouped(sorted_events):
        if kind == "slide":
            label = f"第 {ev.index} 页" + ("（翻回）" if ev.is_backtrack else "")
            h = doc.add_heading(f"[{fmt_ts(ev.ts)}] {label}", level=2)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            opened = True
            if include_images:
                img = session_root / ev.image
                if img.exists():
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(str(img), width=Inches(6.2))
        else:  # para：段落式文本（段首一个时间戳）
            if not opened and not has_slide:
                doc.add_heading("[00:00:00] 开场", level=2)
                opened = True
            para = doc.add_paragraph()
            para.paragraph_format.first_line_indent = Pt(22)  # 首行缩进，段落感更强
            r = para.add_run(f"[{fmt_ts(ev.start_ts)}] ")
            r.bold = True
            para.add_run(ev.text)

    out = session_root / "transcript.docx"
    doc.save(str(out))
    return out
