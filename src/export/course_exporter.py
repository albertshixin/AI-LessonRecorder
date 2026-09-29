# -*- coding: utf-8 -*-
"""AI 教程文档导出：解析教程 Markdown（含 {{slide:N}} 占位符）→ Word 图文文档"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.core.timeline import SlideEvent
from src.utils.logger import logger

_SLIDE_RE = re.compile(r"\{\{slide:(\d+)\}\}")
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _set_cjk_font(doc: Document) -> None:
    from docx.oxml.ns import qn
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def export_course_docx(session_root: Path, course_md_text: str,
                       slide_events: list[SlideEvent]) -> Path:
    """将 AI 生成的教程 MD 转为 docx，占位符替换为真实截图"""
    slide_map: dict[int, str] = {}
    for ev in slide_events:
        slide_map.setdefault(ev.index, ev.image)

    doc = Document()
    _set_cjk_font(doc)

    for raw_line in course_md_text.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if not stripped:
            continue

        # 占位符独立成行 → 插图
        m = _SLIDE_RE.fullmatch(stripped)
        if m:
            idx = int(m.group(1))
            img_rel = slide_map.get(idx)
            if img_rel and (session_root / img_rel).exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(session_root / img_rel), width=Inches(6.2))
            else:
                logger.warning(f"教程插图占位 slide:{idx} 无对应截图，已跳过")
            continue

        # 标题
        if stripped.startswith("#"):
            level = min(len(stripped) - len(stripped.lstrip("#")), 4)
            text = stripped.lstrip("#").strip()
            doc.add_heading(text, level=level)
            continue

        # 表格行（简单处理为文本）
        if stripped.startswith("|"):
            doc.add_paragraph(stripped.replace("|", "  ").strip("-: "))
            continue

        # 列表
        if stripped.startswith(("- ", "* ", "• ")):
            _add_rich_text(doc.add_paragraph(style="List Bullet"), stripped[2:].strip())
            continue
        if re.match(r"^\d+[.、]", stripped):
            _add_rich_text(doc.add_paragraph(style="List Number"),
                           re.sub(r"^\d+[.、]", "", stripped).strip())
            continue

        # 引用
        if stripped.startswith(">"):
            _add_rich_text(doc.add_paragraph(style="Intense Quote"),
                           stripped.lstrip("> ").strip())
            continue

        # 普通段落（可能内嵌占位符 → 行内文字后插图）
        inline = _SLIDE_RE.search(stripped)
        if inline:
            text_part = _SLIDE_RE.sub("", stripped).strip()
            if text_part:
                _add_rich_text(doc.add_paragraph(), text_part)
            idx = int(inline.group(1))
            img_rel = slide_map.get(idx)
            if img_rel and (session_root / img_rel).exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(session_root / img_rel), width=Inches(6.2))
            continue

        _add_rich_text(doc.add_paragraph(), stripped)

    out = session_root / "course_notes.docx"
    doc.save(str(out))
    return out


def _add_rich_text(paragraph, text: str) -> None:
    """支持 **粗体** 的行内富文本"""
    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        r = paragraph.add_run(m.group(1))
        r.bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])
